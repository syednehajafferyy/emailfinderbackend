import csv
import io
import re
import uuid
import shutil
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.config import settings, UPLOADS_DIR, EXPORTS_DIR
from app.models import UploadPreviewResponse, WebsiteResult, ScanStatus
from app.scanner import validate_and_resolve_url

URL_HEADER_PATTERNS = [
    re.compile(r"^(website|url|web|domain|homepage|site|link)$", re.IGNORECASE),
    re.compile(r"(website|web\s*url|domain|homepage|web\s*site)", re.IGNORECASE),
]

COMPANY_HEADER_PATTERNS = [
    re.compile(r"^(company|company\s*name|business|organization|firm|client|brand|name)$", re.IGNORECASE),
    re.compile(r"(company|business|organization|firm|client|brand)", re.IGNORECASE),
]

DOMAIN_REGEX = re.compile(
    r"(https?:\/\/)?([a-zA-Z0-9-]+\.)+[a-zA-Z]{2,}(:\d+)?(\/[^\s]*)?",
    re.IGNORECASE
)

def detect_columns(headers: List[str], sample_rows: List[Dict[str, Any]]) -> Tuple[Optional[str], Optional[str]]:
    """Detects website URL column and company name column from headers and values."""
    detected_url_col: Optional[str] = None
    detected_company_col: Optional[str] = None

    # 1. First attempt by header name exact matches
    for h in headers:
        clean_h = h.strip()
        if not clean_h:
            continue
        if not detected_url_col and URL_HEADER_PATTERNS[0].search(clean_h):
            detected_url_col = clean_h
        elif not detected_company_col and COMPANY_HEADER_PATTERNS[0].search(clean_h):
            detected_company_col = clean_h

    # 2. Second attempt by header name substring matches
    if not detected_url_col:
        for h in headers:
            clean_h = h.strip()
            if URL_HEADER_PATTERNS[1].search(clean_h):
                detected_url_col = clean_h
                break

    if not detected_company_col:
        for h in headers:
            clean_h = h.strip()
            if clean_h != detected_url_col and COMPANY_HEADER_PATTERNS[1].search(clean_h):
                detected_company_col = clean_h
                break

    # 3. Third attempt by inspecting sample values for URL column
    if not detected_url_col and sample_rows:
        col_scores: Dict[str, int] = {h: 0 for h in headers}
        for row in sample_rows:
            for h in headers:
                val = str(row.get(h, "")).strip()
                if not val:
                    continue
                if DOMAIN_REGEX.search(val) or val.startswith("www.") or "://" in val:
                    col_scores[h] += 1
        best_col = max(col_scores.items(), key=lambda x: x[1], default=(None, 0))
        if best_col[1] > 0:
            detected_url_col = best_col[0]

    return detected_url_col, detected_company_col

def parse_uploaded_file(file_path: Path, filename: str) -> Tuple[List[str], List[Dict[str, Any]], List[List[Any]]]:
    """
    Parses an Excel or CSV file.
    Returns: (headers, rows_as_dicts, raw_rows_as_lists)
    """
    ext = file_path.suffix.lower()
    headers: List[str] = []
    rows_as_dicts: List[Dict[str, Any]] = []
    raw_rows: List[List[Any]] = []

    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        sheet = wb.active
        all_rows = list(sheet.iter_rows(values_only=True))
        wb.close()

        if not all_rows:
            return [], [], []

        # Find header row (first non-empty row)
        header_row_idx = 0
        for idx, row in enumerate(all_rows):
            if any(cell is not None and str(cell).strip() for cell in row):
                header_row_idx = idx
                break

        raw_headers = all_rows[header_row_idx]
        headers = [str(h).strip() if h is not None and str(h).strip() else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

        for r_idx, row in enumerate(all_rows[header_row_idx + 1:], start=header_row_idx + 1):
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            row_dict = {}
            row_list = list(row)
            for h_idx, h in enumerate(headers):
                val = row[h_idx] if h_idx < len(row) else None
                row_dict[h] = str(val).strip() if val is not None else ""
            rows_as_dicts.append(row_dict)
            raw_rows.append(row_list)

    elif ext == ".csv":
        # Try UTF-8 first, fallback to latin-1
        content = None
        for enc in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue

        if not content:
            raise ValueError("Unable to read CSV file with standard encodings.")

        reader = csv.reader(io.StringIO(content))
        all_rows = list(reader)
        if not all_rows:
            return [], [], []

        header_row_idx = 0
        for idx, row in enumerate(all_rows):
            if any(cell.strip() for cell in row):
                header_row_idx = idx
                break

        raw_headers = all_rows[header_row_idx]
        headers = [str(h).strip() if str(h).strip() else f"Column_{i+1}" for i, h in enumerate(raw_headers)]

        for row in all_rows[header_row_idx + 1:]:
            if not any(str(cell).strip() for cell in row):
                continue
            row_dict = {}
            for h_idx, h in enumerate(headers):
                val = row[h_idx] if h_idx < len(row) else ""
                row_dict[h] = val.strip()
            rows_as_dicts.append(row_dict)
            raw_rows.append(row)
    else:
        raise ValueError(f"Unsupported file format: {ext}. Please upload a .xlsx or .csv file.")

    return headers, rows_as_dicts, raw_rows

def analyze_spreadsheet(
    file_id: str,
    filename: str,
    file_path: Path
) -> UploadPreviewResponse:
    """Analyzes spreadsheet rows, detects columns, and compiles URL stats."""
    headers, rows_as_dicts, _ = parse_uploaded_file(file_path, filename)
    if not headers or not rows_as_dicts:
        raise ValueError("The uploaded spreadsheet appears to be empty or has no recognizable data rows.")

    sample_preview = rows_as_dicts[:10]
    detected_url_col, detected_company_col = detect_columns(headers, sample_preview)

    total_rows = len(rows_as_dicts)
    empty_urls = 0
    invalid_urls = 0
    valid_urls = 0
    seen_urls: set = set()
    duplicate_urls = 0

    col_to_check = detected_url_col or (headers[0] if headers else None)

    if col_to_check:
        for row in rows_as_dicts:
            val = row.get(col_to_check, "").strip()
            if not val:
                empty_urls += 1
                continue
            is_safe, norm_url, err = validate_and_resolve_url(val)
            if not is_safe:
                invalid_urls += 1
            else:
                if norm_url in seen_urls:
                    duplicate_urls += 1
                else:
                    seen_urls.add(norm_url)
                valid_urls += 1

    return UploadPreviewResponse(
        file_id=file_id,
        filename=filename,
        headers=headers,
        preview_rows=sample_preview,
        detected_url_column=detected_url_col,
        detected_company_column=detected_company_col,
        total_rows=total_rows,
        valid_urls_count=valid_urls,
        empty_urls_count=empty_urls,
        invalid_urls_count=invalid_urls,
        duplicate_urls_count=duplicate_urls,
    )

def export_updated_excel(
    original_file_path: Path,
    original_filename: str,
    job_id: str,
    url_column_name: str,
    results_by_row: Dict[int, WebsiteResult]
) -> Path:
    """
    Generates an updated Excel (.xlsx) file with results inserted immediately after the website column.
    Preserves original row order, formatting, and data. Never overwrites the original file.
    """
    ext = original_file_path.suffix.lower()
    export_filename = f"webmail_finder_results_{job_id}.xlsx"
    export_path = EXPORTS_DIR / export_filename

    new_headers = ["Contact Email", "Source URL", "Status", "Other Emails"]

    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(original_file_path)
        ws = wb.active

        # Find header row and URL column index
        header_row_idx = 1
        url_col_idx = 1
        found_url_col = False

        for r in range(1, min(10, ws.max_row + 1)):
            for c in range(1, ws.max_column + 1):
                cell_val = str(ws.cell(row=r, column=c).value or "").strip()
                if cell_val.lower() == url_column_name.strip().lower():
                    header_row_idx = r
                    url_col_idx = c
                    found_url_col = True
                    break
            if found_url_col:
                break

        # Insert 4 new columns right after the URL column
        insert_at_col = url_col_idx + 1
        ws.insert_cols(insert_at_col, amount=4)

        # Style header cells nicely
        ref_header_cell = ws.cell(row=header_row_idx, column=url_col_idx)
        header_font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        header_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for offset, header_title in enumerate(new_headers):
            cell = ws.cell(row=header_row_idx, column=insert_at_col + offset)
            cell.value = header_title
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        # Populate rows
        # Data rows start after header_row_idx
        data_row_counter = 0
        for current_r in range(header_row_idx + 1, ws.max_row + 1):
            # Check if row is not completely empty
            row_has_data = any(
                ws.cell(row=current_r, column=c).value is not None
                for c in range(1, ws.max_column + 1)
            )
            if not row_has_data:
                continue

            result = results_by_row.get(data_row_counter)
            data_row_counter += 1

            email_val = result.primary_email if result and result.primary_email else ""
            source_val = result.source_url if result and result.source_url else ""
            status_val = result.status.value if result else "Not processed"
            others_val = ", ".join(result.other_emails) if result and result.other_emails else ""

            # Write values
            ws.cell(row=current_r, column=insert_at_col).value = email_val
            ws.cell(row=current_r, column=insert_at_col + 1).value = source_val
            ws.cell(row=current_r, column=insert_at_col + 2).value = status_val
            ws.cell(row=current_r, column=insert_at_col + 3).value = others_val

            for offset in range(4):
                c_cell = ws.cell(row=current_r, column=insert_at_col + offset)
                c_cell.alignment = Alignment(vertical="center")
                c_cell.border = thin_border

        # Auto-adjust column widths for readability
        for offset, header_title in enumerate(new_headers):
            col_letter = openpyxl.utils.get_column_letter(insert_at_col + offset)
            ws.column_dimensions[col_letter].width = max(len(header_title) + 5, 22)

        wb.save(export_path)
        wb.close()

    else:
        # Handle original CSV converted to XLSX export
        headers, rows_as_dicts, _ = parse_uploaded_file(original_file_path, original_filename)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Discovered Emails"

        url_idx = 0
        for i, h in enumerate(headers):
            if h.strip().lower() == url_column_name.strip().lower():
                url_idx = i
                break

        combined_headers = headers[:url_idx + 1] + new_headers + headers[url_idx + 1:]
        ws.append(combined_headers)

        header_font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
        header_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
        for col_num in range(1, len(combined_headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for idx, row in enumerate(rows_as_dicts):
            result = results_by_row.get(idx)
            email_val = result.primary_email if result and result.primary_email else ""
            source_val = result.source_url if result and result.source_url else ""
            status_val = result.status.value if result else "Not processed"
            others_val = ", ".join(result.other_emails) if result and result.other_emails else ""

            row_values = [row.get(h, "") for h in headers]
            combined_row = row_values[:url_idx + 1] + [email_val, source_val, status_val, others_val] + row_values[url_idx + 1:]
            ws.append(combined_row)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

        wb.save(export_path)
        wb.close()

    return export_path

def export_updated_csv(
    original_file_path: Path,
    original_filename: str,
    job_id: str,
    url_column_name: str,
    results_by_row: Dict[int, WebsiteResult]
) -> Path:
    """Generates an updated CSV file with UTF-8 BOM encoding for seamless Excel opening."""
    export_filename = f"webmail_finder_results_{job_id}.csv"
    export_path = EXPORTS_DIR / export_filename

    headers, rows_as_dicts, _ = parse_uploaded_file(original_file_path, original_filename)
    new_headers = ["Contact Email", "Source URL", "Status", "Other Emails"]

    url_idx = 0
    for i, h in enumerate(headers):
        if h.strip().lower() == url_column_name.strip().lower():
            url_idx = i
            break

    combined_headers = headers[:url_idx + 1] + new_headers + headers[url_idx + 1:]

    with open(export_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(combined_headers)

        for idx, row in enumerate(rows_as_dicts):
            result = results_by_row.get(idx)
            email_val = result.primary_email if result and result.primary_email else ""
            source_val = result.source_url if result and result.source_url else ""
            status_val = result.status.value if result else "Not processed"
            others_val = ", ".join(result.other_emails) if result and result.other_emails else ""

            row_values = [row.get(h, "") for h in headers]
            combined_row = row_values[:url_idx + 1] + [email_val, source_val, status_val, others_val] + row_values[url_idx + 1:]
            writer.writerow(combined_row)

    return export_path
